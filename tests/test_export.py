"""Tests for the Excel export helpers (Tier 0 + Tier 2).

neutralize_formula is stdlib-only and always runs. children_to_dataframe
needs pandas and skips cleanly when pandas is not installed.
"""

import os
import tempfile
import unittest

from exporters import neutralize_formula, EXPORT_COLUMNS

try:
    import pandas as pd
    from exporters import children_to_dataframe
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class TestNeutralizeFormula(unittest.TestCase):
    def test_formula_triggers_are_prefixed(self):
        for evil in ("=1+1", "+1+1", "-1+1", "@SUM(A1)", "=HYPERLINK(\"x\")"):
            out = neutralize_formula(evil)
            self.assertIsInstance(out, str)
            self.assertTrue(out.startswith("'"), repr(out))
            self.assertNotEqual(out, evil)
            # the original payload survives as visible text
            self.assertTrue(out[1:].startswith(evil[0]), repr(out))

    def test_ordinary_strings_unchanged(self):
        for ok in ("Иванов Иван", "Младшая группа", "A1", "  = spaced",
                   "x=y", "a+b", "n-m", "user@host", ""):
            self.assertEqual(neutralize_formula(ok), ok)

    def test_non_strings_unchanged(self):
        for v in (123, 12.5, None, True):
            self.assertIs(neutralize_formula(v), v)


@unittest.skipUnless(HAS_PANDAS, "pandas not installed")
class TestChildrenToDataFrame(unittest.TestCase):
    def _rows(self):
        return [
            {"id": 1, "full_name": "Иванов Иван", "group_name": "Младшая"},
            {"id": 2, "full_name": "=группа-уловка", "group_name": "-Старшая"},
        ]

    def test_column_names(self):
        df = children_to_dataframe(self._rows())
        self.assertEqual(list(df.columns), EXPORT_COLUMNS)

    def test_id_column_is_numeric(self):
        df = children_to_dataframe(self._rows())
        self.assertTrue(pd.api.types.is_integer_dtype(df["ID"]))
        self.assertEqual(df["ID"].tolist(), [1, 2])

    def test_text_columns_stay_text_and_cannot_be_formulas(self):
        df = children_to_dataframe(self._rows())
        # pandas 3 stores text in a dedicated ``str`` dtype instead of the old
        # catch-all ``object``; assert the semantic property ("this column holds
        # text") via is_string_dtype so the test holds under both spellings.
        for col in ("ФИО", "Группа"):
            self.assertTrue(
                pd.api.types.is_string_dtype(df[col]),
                f"{col}: dtype {df[col].dtype!r} is not a text dtype",
            )
        self.assertTrue(pd.api.types.is_integer_dtype(df["ID"]))
        # value level: real strings, and none can be parsed as a formula
        for col in ("ФИО", "Группа"):
            for v in df[col]:
                self.assertIsInstance(v, str)
                self.assertFalse(v[:1] in ("=", "+", "-", "@"), (col, v))

    def test_missing_group_becomes_empty_string(self):
        df = children_to_dataframe([{"id": 9, "full_name": "Петров", "group_name": None}])
        self.assertEqual(df["Группа"].iloc[0], "")
        self.assertEqual(df["ФИО"].iloc[0], "Петров")

    def test_empty_rows_produce_empty_frame_with_columns(self):
        df = children_to_dataframe([])
        self.assertEqual(list(df.columns), EXPORT_COLUMNS)
        self.assertEqual(len(df), 0)


@unittest.skipUnless(HAS_PANDAS and HAS_OPENPYXL, "pandas/openpyxl not installed")
class TestExcelRoundTrip(unittest.TestCase):
    """Real XLSX round trip: build via exporters, write with pandas, reopen
    with openpyxl and assert how the cells were actually STORED.

    No tracked database is involved: the rows are synthetic.
    """

    def _rows(self):
        return [
            {"id": 1, "full_name": "Иванов Иван", "group_name": "Младшая"},
            {"id": 2, "full_name": "=группа-уловка", "group_name": "@Старшая"},
            {"id": 3, "full_name": "+имя", "group_name": "-Группа"},
            {"id": 4, "full_name": "Петров Пётр", "group_name": None},
        ]

    def _roundtrip(self):
        df = children_to_dataframe(self._rows())
        tmpdir = tempfile.mkdtemp()
        path = os.path.join(tmpdir, "export.xlsx")
        try:
            df.to_excel(path, index=False)
            self.assertTrue(os.path.isfile(path))
            self.assertGreater(os.path.getsize(path), 0)
            wb = openpyxl.load_workbook(path)
            try:
                return wb.active
            finally:
                wb.close()
        finally:
            import shutil
            shutil.rmtree(tmpdir, ignore_errors=True)

    def test_header_row(self):
        ws = self._roundtrip()
        self.assertEqual(
            [ws.cell(row=1, column=c).value for c in range(1, 4)],
            EXPORT_COLUMNS,
        )

    def test_id_cell_is_numeric(self):
        ws = self._roundtrip()
        for row, expected in ((2, 1), (3, 2), (4, 3), (5, 4)):
            cell = ws.cell(row=row, column=1)
            self.assertEqual(cell.data_type, "n", f"{cell.coordinate} not numeric")
            self.assertEqual(cell.value, expected)
            self.assertIsInstance(cell.value, int)

    def test_text_cells_are_never_formulas(self):
        ws = self._roundtrip()
        # every text cell, including the hostile ones, must be stored as TEXT
        # and never as a formula (data_type 'f'). An empty-string group is
        # read back by openpyxl as None (an empty cell), which is still text,
        # never a formula.
        for row in range(2, ws.max_row + 1):
            for col in (2, 3):
                cell = ws.cell(row=row, column=col)
                self.assertNotEqual(
                    cell.data_type, "f", f"{cell.coordinate} stored as a formula"
                )
                self.assertTrue(
                    cell.value is None or isinstance(cell.value, str),
                    f"{cell.coordinate}: unexpected non-text value {cell.value!r}",
                )

    def test_formula_trigger_strings_stay_text_with_marker(self):
        ws = self._roundtrip()
        # row 3 (id 2): "=группа-уловка" / "@Старшая"
        self.assertEqual(ws.cell(row=3, column=2).value, "'=группа-уловка")
        self.assertEqual(ws.cell(row=3, column=3).value, "'@Старшая")
        # row 4 (id 3): "+имя" / "-Группа"
        self.assertEqual(ws.cell(row=4, column=2).value, "'+имя")
        self.assertEqual(ws.cell(row=4, column=3).value, "'-Группа")
        for row in (3, 4):
            for col in (2, 3):
                self.assertNotEqual(
                    ws.cell(row=row, column=col).data_type, "f"
                )

    def test_cyrillic_roundtrips_intact(self):
        ws = self._roundtrip()
        self.assertEqual(ws.cell(row=2, column=2).value, "Иванов Иван")
        self.assertEqual(ws.cell(row=2, column=3).value, "Младшая")
        self.assertEqual(ws.cell(row=5, column=2).value, "Петров Пётр")

    def test_missing_group_exports_as_empty_string(self):
        ws = self._roundtrip()
        cell = ws.cell(row=5, column=3)
        # The DataFrame holds "" (see TestChildrenToDataFrame); openpyxl maps
        # an empty text cell to None on read-back, so accept both spellings of
        # "no group". What must never happen is a formula or a "None" string.
        self.assertIn(cell.value, (None, ""))
        self.assertNotEqual(cell.data_type, "f")


if __name__ == "__main__":
    unittest.main()
